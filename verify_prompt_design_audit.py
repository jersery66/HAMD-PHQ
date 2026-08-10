#!/usr/bin/env python3
"""Fail-closed QA for the prompt-design audit artifacts."""

from __future__ import annotations

import csv
import hashlib
import json
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_SHEETS = [
    "01_核心说明",
    "02_原始字段盘点",
    "03_唯一提示词模板",
    "04_条件设计字典",
    "05_提示词差分",
    "06_提示词特征",
    "07_人工复核",
    "08_操纵真实性",
    "09_冻结分析输入",
    "10_数据质量问题",
    "11_编码规则",
    "12_结果索引",
]

FEATURE_CSV = "05_prompt_feature_coding.csv"
MANUAL_CSV = "06_prompt_feature_manual_review.csv"
PAIR_CSV = "04_prompt_pairwise_diff.csv"
FROZEN_CSV = "08_prompt_analysis_frozen_input.csv"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main() -> int:
    root = Path(sys.argv[1] if len(sys.argv) > 1 else "prompt_design_audit_20260810").resolve()
    meta = json.loads((root / "audit_metadata.json").read_text(encoding="utf-8"))
    source = Path(meta["input_path"])
    issues: list[str] = []
    checks: dict[str, object] = {}

    source_sha = sha256_file(source)
    checks["source_sha256_unchanged"] = source_sha == meta["input_sha256"] == meta["input_sha256_after_generation"]
    if not checks["source_sha256_unchanged"]:
        issues.append("source workbook SHA-256 changed")

    templates = read_csv(root / "02_unique_prompt_templates.csv")
    template_ids = [r["prompt_template_id"] for r in templates]
    template_hashes = [r["normalized_sha256"] for r in templates]
    checks["template_id_unique"] = len(template_ids) == len(set(template_ids))
    checks["normalized_hash_unique"] = len(template_hashes) == len(set(template_hashes))
    if not checks["template_id_unique"] or not checks["normalized_hash_unique"]:
        issues.append("template IDs or normalized hashes are duplicated")

    frozen = read_csv(root / FROZEN_CSV)
    contexts = {(r["dataset"], r["model"], r["mode"]) for r in frozen}
    c01_c16_counts = Counter(k for k in contexts for r in frozen if (r["dataset"], r["model"], r["mode"]) == k and 1 <= int(r["condition"][1:]) <= 16)
    checks["observed_context_n"] = len(contexts)
    checks["c01_c16_contexts_with_16_conditions"] = sum(v == 16 for v in c01_c16_counts.values())
    checks["c01_c16_contexts_total"] = len(c01_c16_counts)
    if any(v != 16 for v in c01_c16_counts.values()):
        issues.append("at least one observed dataset×model×mode context does not have all C01-C16 rows")

    pairs = read_csv(root / PAIR_CSV)
    pair_groups = Counter((r["dataset"], r["model"], r["mode"]) for r in pairs)
    checks["pair_rows"] = len(pairs)
    checks["pair_group_counts"] = {"|".join(k): v for k, v in pair_groups.items()}
    if any(v != 6 for v in pair_groups.values()):
        issues.append("pairwise diff does not contain exactly six predefined pairs per observed group")

    feature_rows = read_csv(root / FEATURE_CSV)
    manual_rows = read_csv(root / MANUAL_CSV)
    feature_review_keys = set()
    for r in feature_rows:
        for name, value in r.items():
            if name.endswith("_review_required") and value == "yes":
                feature_name = name[: -len("_review_required")]
                feature_review_keys.add((r["dataset"], r["model"], r["mode"], r["condition"], feature_name))
    manual_feature_keys = {
        (r["dataset"], r["model"], r["mode"], r["condition"], r["feature_name"])
        for r in manual_rows
        if r.get("issue_type") == "feature_uncertain_or_low_confidence"
    }
    checks["feature_review_n"] = len(feature_review_keys)
    checks["feature_review_covered_by_manual_review"] = feature_review_keys <= manual_feature_keys
    if not checks["feature_review_covered_by_manual_review"]:
        issues.append("uncertain/low-confidence prompt features are missing from manual review")

    banned_exact = {
        "gold_total",
        "gold_score",
        "gold_level",
        "prediction",
        "predicted_score",
        "mae",
        "nae",
        "correct",
        "sensitivity",
        "specificity",
        "bias",
        "f1",
        "absolute_error",
        "signed_error",
        "accuracy",
        "false_positive",
    }
    frozen_bad_headers = [h for h in frozen[0] if h.lower() in banned_exact] if frozen else []
    checks["frozen_input_has_no_outcome_headers"] = not frozen_bad_headers
    checks["frozen_bad_headers"] = frozen_bad_headers
    if frozen_bad_headers:
        issues.append("frozen input contains outcome headers: " + ",".join(frozen_bad_headers))

    # Same dataset×condition must resolve to one static template; model/mode
    # differences are retained in frozen input and never silently collapsed.
    condition_templates = defaultdict(set)
    for row in frozen:
        condition_templates[(row["dataset"], row["condition"])].add(row["prompt_template_id"])
    template_collisions = {"|".join(k): sorted(v) for k, v in condition_templates.items() if len(v) > 1}
    checks["same_condition_multiple_templates"] = template_collisions
    if template_collisions:
        issues.append("same dataset×condition maps to multiple prompt templates")

    xlsx = next(root.glob("PDCH_PHQ_HAMD_提示词设计与操纵审计总表_*.xlsx"), None)
    checks["xlsx_exists"] = bool(xlsx and xlsx.exists())
    if not xlsx:
        issues.append("audit xlsx is missing")
    else:
        checks["xlsx_sha256"] = sha256_file(xlsx)
        try:
            with zipfile.ZipFile(xlsx) as zf:
                bad = zf.testzip()
                checks["xlsx_zip_valid"] = bad is None
                if bad:
                    issues.append("xlsx zip member failed: " + bad)
            wb = load_workbook(xlsx, read_only=True, data_only=False)
            checks["xlsx_sheet_names"] = list(wb.sheetnames)
            checks["xlsx_sheet_names_exact"] = wb.sheetnames == EXPECTED_SHEETS
            if wb.sheetnames != EXPECTED_SHEETS:
                issues.append("xlsx sheet names/order differ from requested output")
            formula_errors = []
            formula_count = 0
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=False):
                    for cell in row:
                        v = cell.value
                        if isinstance(v, str) and v.startswith("="):
                            formula_count += 1
                        if isinstance(v, str) and v.startswith("#") and v.upper() in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}:
                            formula_errors.append(f"{ws.title}!{cell.coordinate}:{v}")
            checks["xlsx_formula_count"] = formula_count
            checks["xlsx_formula_errors"] = formula_errors
            if formula_errors:
                issues.append("xlsx contains formula errors")
            wb.close()
        except Exception as exc:
            checks["xlsx_open_error"] = repr(exc)
            issues.append("xlsx cannot be opened: " + repr(exc))

    report = {
        "status": "PASS" if not issues else "FAIL",
        "root": str(root),
        "checks": checks,
        "issues": issues,
        "scope_guard": "No prompt-feature to performance association was computed.",
    }
    (root / "prompt_design_audit_qa.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not issues else 1


if __name__ == "__main__":
    raise SystemExit(main())
